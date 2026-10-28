#!/usr/bin/env python3
"""ChambaKit: el archivo de registros de la copia (11/REG-LOCAL).

Uso, con el Python del entorno virtual de la copia (en Windows, .venv\\Scripts\\python.exe):
    .venv/bin/python -I herramientas/registros.py crear
    .venv/bin/python -I herramientas/registros.py agregar [FILAS.json]
    .venv/bin/python -I herramientas/registros.py actualizar [CAMBIOS.json]
    .venv/bin/python -I herramientas/registros.py importar SALIDA.json DECISIONES.json
    .venv/bin/python -I herramientas/registros.py migrar ENTRADA.xlsx [--correcciones CORRECCIONES.json]
    .venv/bin/python -I herramientas/registros.py leer PESTAÑA [--id ID ...] [--columna COLUMNA]
    .venv/bin/python -I herramientas/registros.py informe BUS-ID
    .venv/bin/python -I herramientas/registros.py instantanea

El archivo es <archivo.nombre>.xlsx en la raíz de la copia, con las pestañas del bloque de 11/REG-ESQUEMA y
sus columnas guardadas: las calculadas de informes no se guardan, las da la orden informe.

- crear: el archivo vacío, con los encabezados. No sobrescribe uno que ya exista.
- agregar: filas nuevas, {pestaña: [{columna: texto}]}, de varias pestañas a la vez. Las referencias se
  resuelven contra lo ya registrado más lo de la misma llamada, así que una vacante nueva y su primera
  observación van juntas.
- actualizar: columnas de filas ya registradas, {pestaña: [{columna_id: ID, columna: texto nuevo}]}. El
  ID no cambia; el texto vacío borra el valor.
- importar: una búsqueda desde la salida JSON de buscar.py (07/BUS-REGISTRO), sin pasar sus filas por el agente.
  Crea su fila de busquedas, con la cobertura y la comparación con la última búsqueda comparable; las vacantes
  nuevas, con primera_deteccion = la fecha de la búsqueda; una observación por vacante listada, y actualiza en las
  ya registradas (11/REG-IDENTIDAD) los campos que cambiaron y ultima_observacion_id. Las decisiones son
  {"solicitud_alcance": texto, "recursos_versiones": texto, "no_revisado_pendientes": texto,
   "revisar": {EMP-ID: {Job ID o URL, como en la línea «revisar»: {"region_relacionada": valor, "motivo_region":
   texto, "estado_vigencia": valor, "motivo": texto}}}}; cada línea «revisar ubicación» pide su región y cada
  «revisar vigencia», su estado, y una empresa en [FALLO] se nombra en no_revisado_pendientes. Da el busqueda_id y
  una línea «VAC-ID · título · departamento» por vacante de la búsqueda.
- migrar: pasa a esta versión una hoja de Google descargada como .xlsx o el archivo de registros de una versión
  anterior (ADR 0006). Lee los valores de cada pestaña por el nombre de sus columnas, ignora las calculadas de
  informes y quita las que el esquema ya no tiene (canal), aplica las correcciones, {pestaña: {ID o "fila N":
  {columna: texto}}}, y comprueba cada fila como agregar. Escribe el archivo completo y da, por pestaña, las filas
  de la entrada y las del archivo escrito, más las columnas quitadas, las filas corregidas y las pestañas ignoradas. Solo lee la entrada. No sobrescribe un archivo de registros con filas,
  salvo cuando la entrada es ese mismo archivo.
- leer: las filas de una pestaña, o solo las que tienen esos IDs en la primera columna o en --columna.
- informe: la tabla resumen de una búsqueda (08/EVAL-INFORME), con las columnas calculadas de informes
  (11/REG-CAMPOS) a la fecha de hoy en la zona horaria de 04 y con el umbral de «antigua» de 08/EVAL-ACUERDOS.
- instantanea: regenera completa la instantánea que lee el canal Web (11/REG-INSTANTANEA), en instantanea/:
  informes.md, con las últimas 10 búsquedas y su tabla resumen, y vacantes.md, con una sección por vacante
  evaluada cuya última observación no es Cerrada. Agregar, actualizar, importar y migrar la regeneran después de
  escribir. Si pasa de 300 KB, lo avisa en stderr sin fallar.

Agregar, actualizar, importar y migrar comprueban cada fila que escriben contra el esquema: tipos, listas
(empresa_id, contra 06/CAT-EMPRESAS), obligatoriedad, IDs únicos y referencias. Con un solo error no
escriben nada: dicen todos en stderr y terminan con código 1. El JSON va a stdout. Si falta la
configuración, el archivo u openpyxl, o se pide una pestaña o búsqueda inexistente, termina con código 2
y dice qué falla.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
from zoneinfo import ZoneInfo

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MODULO_04 = os.path.join(RAIZ, "04_PREFERENCIAS_LABORALES.md")
MODULO_06 = os.path.join(RAIZ, "06_EMPRESAS_OBJETIVO.md")
MODULO_08 = os.path.join(RAIZ, "08_EVALUACION_DE_VACANTES.md")
MODULO_11 = os.path.join(RAIZ, "11_ESQUEMAS_DE_REGISTROS.md")
CARPETA_INSTANTANEA = "instantanea"
ULTIMAS_BUSQUEDAS = 10
AVISO_INSTANTANEA_KB = 300
UMBRAL_DE_ANTIGUA = 20


def error(mensaje):
    """Configuración, archivo o biblioteca que falta: termina con código 2 sin tocar el archivo."""
    print(mensaje, file=sys.stderr)
    sys.exit(2)


def leer_modulo(ruta, para):
    try:
        with open(ruta, encoding="utf-8") as f:
            return f.read()
    except OSError as e:
        error(f"{ruta}: no se pudo leer el módulo, de donde salen {para} ({e.strerror}).")


def seccion(md, titulo, nivel="##"):
    """Texto de la sección `<nivel> <titulo>` hasta la siguiente del mismo nivel o mayor."""
    m = re.search(rf"^{nivel} {re.escape(titulo)}(?![\w-]).*?(?=^#{{2,{len(nivel)}}} |\Z)", md, re.M | re.S)
    return m.group(0) if m else None


def bloque_json(texto):
    """Contenido del primer bloque ```json del texto, o None."""
    m = texto and re.search(r"^```json\n(.*?)^```", texto, re.M | re.S)
    return m.group(1) if m else None


def leer_esquema():
    """Bloque de 11/REG-ESQUEMA."""
    ruta = MODULO_11
    bloque = bloque_json(seccion(leer_modulo(ruta, "las pestañas y sus columnas"), "REG-ESQUEMA"))
    if bloque is None:
        error(f"{ruta}: falta el bloque del esquema (```json) en la sección REG-ESQUEMA.")
    try:
        esquema = json.loads(bloque)
    except ValueError as e:
        error(f"{ruta}: el bloque de REG-ESQUEMA no es JSON válido ({e}).")
    if not (isinstance(esquema, dict) and isinstance(esquema.get("archivo"), dict)
            and isinstance(esquema["archivo"].get("nombre"), str) and isinstance(esquema.get("pestanas"), list)):
        error(f"{ruta}, REG-ESQUEMA: el bloque debe ser un objeto con archivo.nombre y pestanas.")
    return esquema


def guardadas(pestana):
    """Las columnas que guarda el archivo: todas menos las calculadas, que son una vista."""
    return [c for c in pestana["columnas"] if c["tipo"] != "calculada"]


def encabezados(pestana):
    """Los nombres de las columnas guardadas, en orden: la fila 1 de la pestaña en el archivo."""
    return [c["nombre"] for c in guardadas(pestana)]


def clave(pestana):
    """La primera columna, la que identifica la fila: su ID, o en informes, su búsqueda."""
    return pestana["columnas"][0]

def pestana_del_esquema(esquema, nombre):
    for p in esquema["pestanas"]:
        if p["nombre"] == nombre:
            return p
    error(f"{nombre!r} no es una pestaña de 11/REG-ESQUEMA; opciones: "
          f"{', '.join(p['nombre'] for p in esquema['pestanas'])}.")


def biblioteca_xlsx():
    """openpyxl, o termina con código 2 si este Python no la tiene."""
    try:
        import openpyxl as biblioteca
    except ImportError:
        error("Falta openpyxl. Corre la herramienta con el Python del entorno virtual de la copia "
              "(.venv/bin/python; en Windows, .venv\\Scripts\\python.exe). Si no existe, lo instala "
              "`/configurar`, paso 8 (Registros).")
    return biblioteca


def ruta_archivo(esquema):
    return os.path.join(RAIZ, esquema["archivo"]["nombre"] + ".xlsx")


def libro_vacio(esquema):
    """El libro con las pestañas del esquema y solo sus encabezados."""
    libro = biblioteca_xlsx().Workbook()
    libro.remove(libro.active)
    for p in esquema["pestanas"]:
        pestana = libro.create_sheet(p["nombre"])
        pestana.append(encabezados(p))
        pestana.freeze_panes = "A2"
    return libro


def crear(esquema):
    ruta = ruta_archivo(esquema)
    if os.path.exists(ruta):
        error(f"{ruta} ya existe; no se sobrescribe.")
    libro_vacio(esquema).save(ruta)
    return {"archivo": ruta, "pestanas": {p["nombre"]: encabezados(p)
                                          for p in esquema["pestanas"]}}


def abrir(esquema):
    """El libro y {pestaña: [fila como {columna: texto}]}, con los encabezados comprobados contra 11."""
    ruta = ruta_archivo(esquema)
    if not os.path.exists(ruta):
        error(f"{ruta}: no existe el archivo de registros; lo crea `/configurar`, paso 8 (Registros).")
    libro = biblioteca_xlsx().load_workbook(ruta)
    filas = {}
    for p in esquema["pestanas"]:
        columnas = encabezados(p)
        if p["nombre"] not in libro.sheetnames:
            error(f"{ruta}: falta la pestaña {p['nombre']} de 11/REG-ESQUEMA.")
        valores = libro[p["nombre"]].iter_rows(values_only=True)
        en_el_archivo = list(next(valores, ()))
        if en_el_archivo != columnas:
            error(f"{ruta}, pestaña {p['nombre']}: los encabezados {en_el_archivo} no coinciden con los de "
                  f"11/REG-ESQUEMA, {columnas} (11/REG-DECISIONES, «Encabezados»).")
        filas[p["nombre"]] = [dict(zip(columnas, ("" if v is None else str(v) for v in fila)))
                              for fila in valores if any(v not in (None, "") for v in fila)]
    return libro, filas


def escribir(celda, valor):
    """Un valor como texto. openpyxl tomaría por fórmula un texto que empieza con «=», como uno copiado de un
    anuncio, y una hoja de cálculo lo evaluaría al abrir el archivo."""
    celda.value, celda.number_format = valor or None, "@"
    if valor:
        celda.data_type = "s"


def guardar(libro, ruta):
    """Escribe el libro completo en un archivo temporal y lo pone en lugar del anterior: nunca queda a medias."""
    temporal = ruta + ".tmp"
    try:
        libro.save(temporal)
        os.replace(temporal, ruta)
    finally:
        if os.path.exists(temporal):
            os.remove(temporal)


def rechazar(errores):
    """Filas que no cumplen el esquema: las dice todas y termina con código 1 sin escribir ninguna."""
    for e in errores:
        print(e, file=sys.stderr)
    print("No se escribió ninguna fila.", file=sys.stderr)
    sys.exit(1)


def leer_entrada(archivo):
    """El JSON de las filas, de un archivo o de stdin ("-")."""
    try:
        if archivo == "-":
            return json.load(sys.stdin)
        with open(archivo, encoding="utf-8") as f:
            return json.load(f)
    except OSError as e:
        rechazar([f"{archivo}: no se pudo leer ({e.strerror})."])
    except ValueError as e:
        rechazar([f"{'La entrada' if archivo == '-' else archivo}: no es JSON válido ({e})."])


def filas_de_entrada(esquema, entrada):
    """{pestaña: [{columna: texto}]} de la entrada, con null como texto vacío, o la rechaza si no tiene esa
    forma o nombra pestañas o columnas que el archivo no guarda."""
    if not isinstance(entrada, dict):
        rechazar(["La entrada debe ser un objeto JSON {pestaña: [fila, ...]}."])
    pestanas = {p["nombre"]: p for p in esquema["pestanas"]}
    errores, filas = [], {}
    for nombre, nuevas in entrada.items():
        if nombre not in pestanas:
            errores.append(f"{nombre!r} no es una pestaña de 11/REG-ESQUEMA; opciones: {', '.join(pestanas)}.")
            continue
        if not (isinstance(nuevas, list) and all(isinstance(f, dict) for f in nuevas)):
            errores.append(f"{nombre}: las filas deben ser una lista de objetos {{columna: valor}}.")
            continue
        columnas = {c["nombre"]: c for c in pestanas[nombre]["columnas"]}
        filas[nombre] = []
        for n, fila in enumerate(nuevas, 1):
            for c, v in fila.items():
                if c not in columnas:
                    errores.append(f"{nombre}, fila {n}: {c!r} no es una columna de {nombre}.")
                elif columnas[c]["tipo"] == "calculada":
                    errores.append(f"{nombre}, fila {n}: {c} se calcula (orden informe); no se escribe.")
                elif not isinstance(v, (str, type(None))):
                    errores.append(f"{nombre}, fila {n}: {c} debe ser texto, entre comillas; dice {v!r}.")
            filas[nombre].append({c: v or "" for c, v in fila.items()})
    if errores:
        rechazar(errores)
    return filas


def leer_catalogo():
    """Los EMP-ID de 06/CAT-EMPRESAS."""
    ruta = MODULO_06
    catalogo = seccion(leer_modulo(ruta, "los IDs de empresa_id"), "CAT-EMPRESAS")
    if catalogo is None:
        error(f"{ruta}: falta la sección CAT-EMPRESAS, el catálogo de empresas.")
    return re.findall(r"^\|\s*(EMP-[A-Z0-9]+(?:-[A-Z0-9]+)*)\s*\|", catalogo, re.M)


def es_fecha(valor, con_hora):
    m = re.fullmatch(r"(\d{4}-\d{2}-\d{2})( \d{2}:\d{2})?", valor)
    if not m or (m.group(2) and not con_hora):
        return False
    try:
        dt.datetime.strptime(valor, "%Y-%m-%d %H:%M" if m.group(2) else "%Y-%m-%d")
    except ValueError:
        return False
    return True


def comprobar(esquema, todas, revisar):
    """Errores de las filas de revisar [(pestaña, fila, etiqueta)] según 11/REG-ESQUEMA. Las referencias
    y la unicidad se resuelven contra todas: las ya registradas más las de esta llamada."""
    pestanas = {p["nombre"]: p for p in esquema["pestanas"]}
    por_id = {nombre: {} for nombre in todas}
    for nombre, filas in todas.items():
        primera = clave(pestanas[nombre])["nombre"]
        for f in filas:
            por_id[nombre].setdefault(f[primera], []).append(f)
    catalogo = []
    errores = []
    for nombre, fila, etiqueta in revisar:
        def mal(columna, mensaje):
            errores.append(f"{nombre}, {etiqueta}: {columna}: {mensaje}")

        for c in guardadas(pestanas[nombre]):
            col, v = c["nombre"], fila[c["nombre"]]
            if v == "":
                if c.get("oblig") == "si" and not (c.get("vacia_con") and fila[c["vacia_con"]]):
                    mal(col, "es obligatoria y está vacía.")
                continue
            tipo = c["tipo"]
            if tipo in ("id", "ref") and not re.fullmatch(f"(PRUEBA-)?{c['patron']}", v):
                mal(col, f"«{v}» no tiene la forma {c['patron']}.")
            elif tipo == "id":
                if sum(f[col] == v for f in todas[nombre]) > 1:
                    mal(col, f"«{v}» se repite en {nombre}; los IDs no se reutilizan.")
            elif tipo == "ref":
                citadas = por_id[c["en"]].get(v, [])
                if not citadas:
                    mal(col, f"«{v}» no existe en {c['en']}.")
                elif c.get("misma_vacante") and not any(f["vacante_id"] == fila["vacante_id"] for f in citadas):
                    mal(col, f"«{v}» es de otra vacante, no de {fila['vacante_id']}.")
                if c.get("unica_con") and sum(f[col] == v and f[c["unica_con"]] == fila[c["unica_con"]]
                                              for f in todas[nombre]) > 1:
                    mal(col, f"la pareja {c['unica_con']} «{fila[c['unica_con']]}» + «{v}» se repite en {nombre}.")
                if c.get("distinta_de") and v == fila[c["distinta_de"]]:
                    mal(col, f"no puede ser igual a {c['distinta_de']}.")
            elif tipo == "fecha" and not es_fecha(v, c.get("hora")):
                forma = "AAAA-MM-DD o AAAA-MM-DD HH:MM" if c.get("hora") else "AAAA-MM-DD, sin hora"
                mal(col, f"«{v}» no es una fecha real con la forma {forma}.")
            elif tipo == "url" and not re.fullmatch(r"https?://\S+", v):
                mal(col, f"«{v}» no es una URL que empiece con http:// o https:// y sin espacios.")
            elif tipo == "patron" and not re.fullmatch(c["patron"], v):
                mal(col, f"«{v}» no tiene la forma {c['patron']}.")
            elif tipo == "lista":
                if "valores_de" in c:
                    if c["valores_de"] != "06/CAT-EMPRESAS":
                        error(f"{MODULO_11}, REG-ESQUEMA: valores_de {c['valores_de']!r} desconocido.")
                    catalogo = catalogo or leer_catalogo()
                    if v not in catalogo:
                        mal(col, f"«{v}» no está en el catálogo de 06 (CAT-EMPRESAS); dala de alta ahí primero.")
                elif v not in c["valores"]:
                    mal(col, f"«{v}» no es uno de: {' / '.join(c['valores'])}.")
    return errores


def agregar(esquema, entrada):
    columnas = {p["nombre"]: encabezados(p) for p in esquema["pestanas"]}
    nuevas = {nombre: [{c: f.get(c, "") for c in columnas[nombre]} for f in filas]
              for nombre, filas in filas_de_entrada(esquema, entrada).items()}
    libro, registradas = abrir(esquema)
    todas = {nombre: filas + nuevas.get(nombre, []) for nombre, filas in registradas.items()}
    revisar = [(nombre, fila, f"fila {n}")
               for nombre, filas in nuevas.items() for n, fila in enumerate(filas, 1)]
    errores = comprobar(esquema, todas, revisar)
    if errores:
        rechazar(errores)
    vistas = instantanea(esquema, todas)
    anexar(libro, nuevas)
    guardar(libro, ruta_archivo(esquema))
    guardar_instantanea(vistas)
    return {"agregadas": {nombre: len(filas) for nombre, filas in nuevas.items()}}


def anexar(libro, nuevas):
    """Escribe en el libro las filas nuevas, {pestaña: [fila con todas sus columnas, en orden]}, al final."""
    for nombre, filas in nuevas.items():
        pestana = libro[nombre]
        for fila in filas:
            numero = pestana.max_row + 1
            for i, valor in enumerate(fila.values(), 1):
                escribir(pestana.cell(numero, i), valor)


def escribir_cambios(esquema, libro, cambios):
    """Escribe en el libro las columnas dadas, {pestaña: [{columna_id: ID, columna: valor}]}, en la fila de su ID."""
    for nombre, filas in cambios.items():
        pestana = libro[nombre]
        columnas = encabezados(pestana_del_esquema(esquema, nombre))
        for cambio in filas:
            numero = next(celda.row for celda in pestana["A"] if celda.value == cambio[columnas[0]])
            for columna, valor in cambio.items():
                escribir(pestana.cell(numero, columnas.index(columna) + 1), valor)


def actualizar(esquema, entrada):
    """Cambia las columnas dadas de filas ya registradas, cada una identificada por su ID (la primera columna).
    Las filas cambiadas se comprueban completas, como al agregarlas."""
    cambios = filas_de_entrada(esquema, entrada)
    libro, registradas = abrir(esquema)
    todas = {nombre: [dict(f) for f in filas] for nombre, filas in registradas.items()}
    errores, revisar = [], []
    for nombre, filas in cambios.items():
        id_ = clave(pestana_del_esquema(esquema, nombre))
        if id_["tipo"] != "id":
            errores.append(f"{nombre}: sus filas no tienen ID propio y no se actualizan.")
            continue
        for n, cambio in enumerate(filas, 1):
            valor = cambio.get(id_["nombre"])
            fila = next((f for f in todas[nombre] if valor and f[id_["nombre"]] == valor), None)
            if fila is None:
                errores.append(f"{nombre}, fila {n}: {id_['nombre']} «{valor or ''}» no está registrado; cada "
                               f"cambio lleva el ID de una fila de {nombre}.")
                continue
            fila.update(cambio)
            revisar.append((nombre, fila, valor))
    errores += comprobar(esquema, todas, revisar)
    if errores:
        rechazar(errores)
    vistas = instantanea(esquema, todas)
    escribir_cambios(esquema, libro, cambios)
    guardar(libro, ruta_archivo(esquema))
    guardar_instantanea(vistas)
    return {"actualizadas": {nombre: len(filas) for nombre, filas in cambios.items()}}


# --- importar: la salida de buscar.py (07/BUS-REGISTRO) ------------------------

# Las decisiones de una línea «revisar» de una vacante: cada columna que se decide, con su motivo.
DECIDIBLES = {"region_relacionada": "motivo_region", "estado_vigencia": "motivo"}
DE_LAS_DECISIONES = ("solicitud_alcance", "recursos_versiones", "no_revisado_pendientes")
# Las columnas de vacantes que vienen de buscar.py; en una ya registrada se actualizan si cambiaron.
DE_LA_SALIDA = ("job_id", "titulo", "url_oficial", "ubicacion_literal", "modalidad_literal", "region_relacionada",
                "fecha_publicacion")
# Los conteos de un listado, que cambian entre búsquedas comparables (buscar.py, listado() y cada adaptador).
CUENTAS = re.compile(r"\d+ de \d+ grupos|\d+ únicos(?: de \w+|, sin total publicado)?")


def generador_de_ids(columna, filas):
    """Los IDs siguientes al mayor registrado, con el prefijo y las cifras del patrón (11/REG-CONVENCIONES). Los
    sintéticos, con PRUEBA- delante, no cuentan."""
    m = re.fullmatch(r"([A-Z]+-)\\d\{(\d+),\}", columna["patron"])
    prefijo, cifras = m.group(1), int(m.group(2))
    numeros = [int(f[columna["nombre"]][len(prefijo):]) for f in filas
               if re.fullmatch(re.escape(prefijo) + r"\d+", f[columna["nombre"]])]
    n = max(numeros, default=0)
    while True:
        n += 1
        yield f"{prefijo}{n:0{cifras}d}"


def con_job_id(job_id):
    return job_id != "" and job_id.casefold() != "no publicado"


def identidades(vacante):
    """Las claves de una vacante (11/REG-IDENTIDAD): empresa + Job ID y empresa + URL oficial. La primera es su
    identidad: la de Job ID si lo tiene; si no, la provisional, por la URL."""
    claves = []
    if con_job_id(vacante["job_id"]):
        claves.append((vacante["empresa_id"], "job_id", vacante["job_id"]))
    if vacante["url_oficial"]:
        claves.append((vacante["empresa_id"], "url_oficial", vacante["url_oficial"]))
    return claves


CAMPOS_DE_EMPRESA = ("acceso", "consultas", "total_reportado", "unicos", "completa", "vacantes",
                     "remotas_fuera_de_region", "remotas_revision", "revisar")
CAMPOS_DE_VACANTE = ("job_id", "titulo", "url_oficial", "ubicacion_literal", "modalidad_literal", "region_relacionada",
                     "ubicaciones_por_revisar", "fecha_publicacion", "estado_vigencia", "motivo", "url_fuente",
                     "departamento", "descripcion")


def forma_de_salida(salida):
    """Los errores de forma de la salida de buscar.py: cada empresa con su EMP-ID y su error o sus listados, y cada
    vacante con los campos de vacantes y observaciones (07/BUS-ACCESO-CODE, «Salida»)."""
    if not (isinstance(salida, dict) and isinstance(salida.get("script"), str)
            and isinstance(salida.get("fecha_hora"), str) and isinstance(salida.get("empresas"), list)):
        return ["La salida no tiene la forma de la de buscar.py: un objeto con script, fecha_hora y empresas."]
    errores = []
    for n, r in enumerate(salida["empresas"], 1):
        if not (isinstance(r, dict) and isinstance(r.get("empresa_id"), str)):
            errores.append(f"salida, empresa {n}: no es un objeto con empresa_id.")
        elif "error" not in r:
            faltan = [c for c in CAMPOS_DE_EMPRESA if c not in r]
            faltan += sorted({c for lista in ("vacantes", "remotas_fuera_de_region")
                              for v in r.get(lista) or [] for c in CAMPOS_DE_VACANTE if not isinstance(v, dict)
                              or c not in v})
            if faltan:
                errores.append(f"salida, {r['empresa_id']}: le falta {', '.join(dict.fromkeys(faltan))}; ¿es la salida "
                               "de herramientas/buscar.py sin editar?")
    return errores


def nombre_en_buscar(v):
    """Cómo nombra buscar.py una vacante en sus líneas «revisar»: por su Job ID o, sin él, por su URL."""
    return v["job_id"] if v["job_id"] not in (None, "No publicado") else v["url_oficial"] or v["url_fuente"]


def listadas(r):
    """{nombre en buscar.py: (vacante, sus líneas «revisar»)} de una empresa de la salida: las de la región y las
    remotas de fuera. Las publicaciones de una misma requisición (mismo nombre) son una sola vacante, con todas sus
    ubicaciones y las líneas «revisar» de todas."""
    vacantes = {}
    for v in r["vacantes"] + r["remotas_fuera_de_region"]:
        nombre = nombre_en_buscar(v)
        v = {c: ("" if x is None else x) for c, x in v.items()}
        if nombre not in vacantes:
            vacantes[nombre] = (v, por_revisar(v))
            continue
        primera, lineas = vacantes[nombre]
        ubicaciones = primera["ubicacion_literal"].split(" | ") + v["ubicacion_literal"].split(" | ")
        primera["ubicacion_literal"] = " | ".join(dict.fromkeys(u for u in ubicaciones if u))
        for columna, linea in por_revisar(v).items():
            lineas[columna] = f"{lineas[columna]} · {linea}" if columna in lineas else linea
    return vacantes


def por_revisar(v):
    """Las líneas «revisar» de una vacante que piden decisión, {columna decidible: línea}, como las da buscar.py."""
    lineas = {}
    if v["ubicaciones_por_revisar"] or v["region_relacionada"] == "Indeterminada" or v.get("revisar"):
        notas = [v["ubicacion_literal"]] + list(v.get("revisar") or [])
        lineas["region_relacionada"] = "revisar ubicación: " + "; ".join(n for n in notas if n)
    if v["estado_vigencia"] != "Confirmada abierta" or "detalle no consultado" in v["motivo"]:
        lineas["estado_vigencia"] = f"revisar vigencia: {v['motivo']}"
    return lineas


def fuente(r):
    """El elemento de una empresa en fuentes_consultas_filtros: acceso, listados con sus conteos y remotas."""
    partes = [r["acceso"], *r["consultas"]] + ([f"remotas: {r['remotas_revision']}"] if r["remotas_revision"] else [])
    return f"{r['empresa_id']}: " + "; ".join(partes).replace(" | ", " / ")


def cobertura(r, vacantes):
    """El elemento de una empresa en cobertura_limitaciones (07/BUS-REGISTRO)."""
    if r["completa"] is None:
        texto = (f"{r['unicos']} únicos, sin total publicado; "
                 + (r.get("verificacion_cobertura") or "sin comprobación de la ficha"))
    else:
        texto = f"{r['unicos']} únicos de {r['total_reportado']}, {'completa' if r['completa'] else 'incompleta'}"
    texto += f"; detalle leído en {sum(bool(v['descripcion']) for v in vacantes)} de {len(vacantes)}"
    return f"{r['empresa_id']}: {texto}" + "".join(f"; revisar: {nota}" for nota in r["revisar"])


def comparacion(registradas, r, actuales, cambios):
    """El elemento de una empresa en comparacion_anterior, contra la última búsqueda comparable (07/BUS-COBERTURA):
    la que listó la misma empresa con el mismo acceso y las mismas consultas, salvo los conteos. None si no hay."""
    firma = CUENTAS.sub("#", fuente(r))
    comparables = [b for b in registradas["busquedas"]
                   if any(CUENTAS.sub("#", e) == firma for e in b["fuentes_consultas_filtros"].split(" | "))]
    if not comparables:
        return None
    anterior = max(comparables, key=lambda b: (b["fecha_hora"], b["busqueda_id"]))["busqueda_id"]
    de_la_empresa = {v["vacante_id"] for v in registradas["vacantes"] if v["empresa_id"] == r["empresa_id"]}
    # Una vacante que esa búsqueda vio cerrada (07/BUS-PASOS, paso 7) ya no se espera en el listado.
    previas = list(dict.fromkeys(o["vacante_id"] for o in registradas["observaciones"]
                                 if o["busqueda_id"] == anterior and o["vacante_id"] in de_la_empresa
                                 and o["estado_vigencia"] != "Cerrada"))

    def lista(ids):
        return ", ".join(ids) or "ninguno"

    texto = (f"{r['empresa_id']}, contra {anterior}: nuevos {lista([v for v in actuales if v not in previas])}; "
             f"ya conocidos {lista([v for v in actuales if v in previas])}; no observados esta vez, sin declararlos "
             f"cerrados: {lista([v for v in previas if v not in actuales])}")
    if cambios:
        texto += "; cambios: " + ", ".join(f"{v} ({', '.join(c)})" for v, c in cambios.items())
    return texto


def decisiones_de(decididas, nombre, lineas, etiqueta, errores):
    """{columna: (valor, motivo)} de lo que se decidió para una vacante, comprobado contra sus líneas «revisar»."""
    d = decididas.get(nombre, {})
    if not isinstance(d, dict):
        errores.append(f"decisiones, {etiqueta}: debe ser un objeto {{columna: texto}}.")
        return {}
    for c, v in d.items():
        if c not in DECIDIBLES and c not in DECIDIBLES.values():
            errores.append(f"decisiones, {etiqueta}: {c!r} no se decide; opciones: "
                           + ", ".join(f"{c} con {m}" for c, m in DECIDIBLES.items()) + ".")
        elif not isinstance(v, str):
            errores.append(f"decisiones, {etiqueta}: {c} debe ser texto, entre comillas; dice {v!r}.")
    elegidas = {}
    for columna, motivo in DECIDIBLES.items():
        valor, porque = d.get(columna), d.get(motivo)
        if columna in lineas and not (valor or porque):
            errores.append(f"{etiqueta}: falta la decisión de la línea «{lineas[columna]}»: {columna} y {motivo}.")
        elif bool(valor) != bool(porque):
            errores.append(f"decisiones, {etiqueta}: {columna} va con su {motivo}; falta "
                           f"{motivo if valor else columna}.")
        if isinstance(valor, str) and valor:
            elegidas[columna] = (valor, porque if isinstance(porque, str) else "")
    return elegidas


def importar(esquema, salida, decisiones):
    """Registra una búsqueda desde la salida de buscar.py y las decisiones del agente (07/BUS-REGISTRO): su fila de
    busquedas, las vacantes nuevas, una observación por vacante listada y los cambios a las ya registradas. Con un
    solo error no escribe nada."""
    forma = forma_de_salida(salida)
    if forma:
        rechazar(forma)
    if not isinstance(decisiones, dict):
        rechazar(["Las decisiones deben ser un objeto JSON con " + ", ".join(DE_LAS_DECISIONES) + " y revisar."])
    errores = [f"decisiones: {c!r} no es un campo de las decisiones; opciones: {', '.join(DE_LAS_DECISIONES)}, revisar."
               for c in decisiones if c not in (*DE_LAS_DECISIONES, "revisar")]
    errores += [f"decisiones: {c} debe ser texto, entre comillas; dice {decisiones[c]!r}."
                for c in DE_LAS_DECISIONES if not isinstance(decisiones.get(c, ""), str)]
    revisar = decisiones.get("revisar", {})
    if not (isinstance(revisar, dict) and all(isinstance(d, dict) for d in revisar.values())):
        errores.append("decisiones: revisar debe ser {EMP-ID: {nombre de la vacante en buscar.py: decisión}}.")
        revisar = {}
    pendientes = decisiones.get("no_revisado_pendientes")
    pendientes = pendientes if isinstance(pendientes, str) else ""

    libro, registradas = abrir(esquema)
    todas = {nombre: [dict(f) for f in filas] for nombre, filas in registradas.items()}
    pestanas = {p["nombre"]: p for p in esquema["pestanas"]}
    columnas = {nombre: encabezados(p) for nombre, p in pestanas.items()}
    ids = {nombre: generador_de_ids(clave(pestanas[nombre]), todas[nombre])
           for nombre in ("busquedas", "vacantes", "observaciones")}
    busqueda_id, fecha = next(ids["busquedas"]), salida["fecha_hora"][:10]
    por_id = {v["vacante_id"]: v for v in todas["vacantes"]}
    por_identidad = {}
    for v in todas["vacantes"]:
        for k in identidades(v):
            por_identidad.setdefault(k, v)

    def canonica(v):
        """La fila canónica de una registrada: la que la absorbió, si fue un duplicado (11/REG-IDENTIDAD)."""
        vistas = set()
        while v and v["duplicado_de"] and v["vacante_id"] not in vistas:
            vistas.add(v["vacante_id"])
            v = por_id.get(v["duplicado_de"], v)
        return v

    fuentes, coberturas, comparaciones, compacta = [salida["script"]], [], [], []
    nuevas = {nombre: [] for nombre in ("busquedas", "vacantes", "observaciones")}
    cambios, revisar_filas = {}, []
    for r in salida["empresas"]:
        eid = r["empresa_id"]
        if "error" in r:
            fuentes.append(f"{eid}: [FALLO] {r['error']}")
            coberturas.append(f"{eid}: [FALLO] {r['error']}")
            if not re.search(rf"(?<![\w-]){re.escape(eid)}(?![\w-])", pendientes):
                errores.append(f"{eid}: salió como [FALLO] ({r['error']}) y no está declarada en "
                               "no_revisado_pendientes de las decisiones.")
            continue
        vacantes = listadas(r)
        decididas = revisar.get(eid, {})
        fuentes.append(fuente(r))
        coberturas.append(cobertura(r, [v for v, _ in vacantes.values()]))
        actuales, cambiadas = [], {}
        for nombre, (v, lineas) in vacantes.items():
            etiqueta = f"{eid}, {nombre}"
            elegidas = decisiones_de(decididas, nombre, lineas, etiqueta, errores)
            fila = {c: v[c] for c in DE_LA_SALIDA} | {"empresa_id": eid}
            motivo = v["motivo"]
            estado = v["estado_vigencia"]
            if "estado_vigencia" in elegidas:
                estado, motivo = elegidas["estado_vigencia"][0], elegidas["estado_vigencia"][1] or motivo
            if "region_relacionada" in elegidas:
                fila["region_relacionada"], porque = elegidas["region_relacionada"]
                motivo += f"; región {fila['region_relacionada']}: {porque}"
            observacion_id = next(ids["observaciones"])
            identidad = identidades(fila)[:1]
            registrada = canonica(por_identidad[identidad[0]]) if identidad and identidad[0] in por_identidad else None
            if registrada is None:
                vacante_id = next(ids["vacantes"])
                nuevas["vacantes"].append({c: "" for c in columnas["vacantes"]} | fila | {
                    "vacante_id": vacante_id, "primera_deteccion": fecha, "ultima_observacion_id": observacion_id})
                revisar_filas.append(("vacantes", nuevas["vacantes"][-1], etiqueta))
            else:
                vacante_id = registrada["vacante_id"]
                distintas = {c: fila[c] for c in DE_LA_SALIDA if fila[c] and fila[c] != registrada[c]
                             and (c != "job_id" or not con_job_id(registrada["job_id"]))}
                motivo += "".join(f"; {c}: «{registrada[c]}» → «{x}»" for c, x in distintas.items() if registrada[c])
                if distintas:
                    cambiadas[vacante_id] = list(distintas)
                cambio = {"vacante_id": vacante_id, **distintas, "ultima_observacion_id": observacion_id}
                registrada.update(cambio)
                cambios.setdefault(vacante_id, {}).update(cambio)
                revisar_filas.append(("vacantes", registrada, etiqueta))
            nuevas["observaciones"].append({
                "observacion_id": observacion_id, "vacante_id": vacante_id, "busqueda_id": busqueda_id,
                "fecha_verificacion": fecha, "estado_vigencia": estado, "motivo": motivo,
                "url_fuente": v["url_fuente"], "tipo_fuente": "Oficial"})
            revisar_filas.append(("observaciones", nuevas["observaciones"][-1], etiqueta))
            actuales.append(vacante_id)
            compacta.append(f"{vacante_id} · {fila['titulo']} · {v['departamento'] or '—'}")
        comparaciones.append(comparacion(registradas, r, actuales, cambiadas) or f"{eid}: Sin búsqueda comparable")
        errores += [f"decisiones, {eid}, {nombre}: no es una vacante de {eid} en la salida; se nombra como en las "
                    "líneas «revisar» de buscar.py, por su Job ID o su URL." for nombre in decididas
                    if nombre not in vacantes]
    con_resultados = {r["empresa_id"] for r in salida["empresas"] if "error" not in r}
    errores += [f"decisiones, {eid}: no es una empresa con resultados en la salida." for eid in revisar
                if eid not in con_resultados]
    busqueda = {c: decisiones.get(c, "") if isinstance(decisiones.get(c, ""), str) else "" for c in DE_LAS_DECISIONES}
    busqueda.update({
        "busqueda_id": busqueda_id, "fecha_hora": salida["fecha_hora"],
        "fuentes_consultas_filtros": " | ".join(fuentes), "cobertura_limitaciones": " | ".join(coberturas),
        "comparacion_anterior": " | ".join(comparaciones) if any(c for c in comparaciones
                                                                 if not c.endswith(": Sin búsqueda comparable"))
                                else "Sin búsqueda comparable"})
    nuevas["busquedas"].append({c: busqueda[c] for c in columnas["busquedas"]})
    errores += [f"Esta salida de buscar.py ya se importó como {b['busqueda_id']}: misma fecha_hora y mismas fuentes."
                for b in registradas["busquedas"] if (b["fecha_hora"], b["fuentes_consultas_filtros"])
                == (busqueda["fecha_hora"], busqueda["fuentes_consultas_filtros"])]
    for nombre, filas in nuevas.items():
        todas[nombre] += filas
    errores += comprobar(esquema, todas, [("busquedas", nuevas["busquedas"][0], busqueda_id)] + revisar_filas)
    if errores:
        rechazar(errores)
    vistas = instantanea(esquema, todas)
    anexar(libro, nuevas)
    escribir_cambios(esquema, libro, {"vacantes": list(cambios.values())})
    guardar(libro, ruta_archivo(esquema))
    guardar_instantanea(vistas)
    return {"busqueda_id": busqueda_id, "vacantes": compacta}


# --- migrar: la hoja descargada o un archivo de registros 1.x (ADR 0006) ----------

# Las columnas que el esquema ya no tiene, por pestaña: migrar las quita.
RETIRADAS = {"busquedas": ("canal",)}


def texto_de_celda(valor):
    """Una celda de la entrada como el texto que guarda el archivo. La hoja guardaba texto, pero una celda editada a
    mano puede llegar como fecha o número."""
    if valor is None:
        return ""
    if isinstance(valor, dt.datetime):
        return f"{valor:%Y-%m-%d}" if valor.time() == dt.time() else f"{valor:%Y-%m-%d %H:%M}"
    if isinstance(valor, dt.date):
        return valor.isoformat()
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def tiene_filas(ruta):
    """Si alguna pestaña del libro tiene algo debajo de su fila 1."""
    try:
        libro = biblioteca_xlsx().load_workbook(ruta, read_only=True)
    except Exception as e:
        error(f"{ruta}: no se pudo abrir para ver si tiene filas ({e}); migrar no lo sobrescribe.")
    try:
        return any(any(v not in (None, "") for v in fila)
                   for pestana in libro.worksheets for fila in pestana.iter_rows(min_row=2, values_only=True))
    finally:
        libro.close()


def encabezado(celda, formula):
    """El nombre de una columna de la entrada. En la hoja, el de cada calculada de informes era una fórmula,
    ={"nombre";MAP(…)}: si el archivo no guarda su valor, el nombre sale de la fórmula."""
    m = celda is None and isinstance(formula, str) and re.match(r'=\{"(\w+)";', formula)
    return m.group(1) if m else texto_de_celda(celda).strip()


def leer_para_migrar(esquema, ruta):
    """Las filas de la entrada, {pestaña: [(fila en la hoja, {columna guardada: texto})]}, con las columnas quitadas
    de cada pestaña y las pestañas que el esquema no tiene. De la hoja descargada se leen los valores, no las
    fórmulas. Rechaza la entrada si le falta una pestaña o una columna, o trae una que el esquema no conoce."""
    try:
        libro = biblioteca_xlsx().load_workbook(ruta, data_only=True)
        formulas = biblioteca_xlsx().load_workbook(ruta, read_only=True)
    except OSError as e:
        error(f"{ruta}: no se pudo leer ({e.strerror}).")
    except Exception as e:
        error(f"{ruta}: no se pudo abrir como .xlsx ({e}). Descarga la hoja con Archivo → Descargar → Microsoft "
              "Excel (.xlsx).")
    from openpyxl.utils import get_column_letter
    errores, filas, quitadas = [], {}, {}
    for p in esquema["pestanas"]:
        nombre, columnas = p["nombre"], encabezados(p)
        if nombre not in libro.sheetnames:
            errores.append(f"{nombre}: falta la pestaña en {ruta}.")
            continue
        valores = libro[nombre].iter_rows(values_only=True)
        fila_1 = next(formulas[nombre].iter_rows(max_row=1, values_only=True), ())
        en_la_entrada = [encabezado(v, f) for v, f in zip(next(valores, ()), fila_1 + (None,) * len(columnas))]
        se_quitan = [c["nombre"] for c in p["columnas"] if c["nombre"] not in columnas] + list(RETIRADAS.get(nombre, ()))
        errores += [f"{nombre}: falta la columna {c}." for c in columnas if c not in en_la_entrada]
        errores += [f"{nombre}: la columna {c} se repite." for c in dict.fromkeys(en_la_entrada)
                    if c and en_la_entrada.count(c) > 1]
        errores += [f"{nombre}: {c!r} no es una columna de 11/REG-ESQUEMA ni de las que se quitan "
                    f"({', '.join(se_quitan) or 'ninguna'})." for c in en_la_entrada
                    if c and c not in columnas and c not in se_quitan]
        posicion = {c: en_la_entrada.index(c) for c in columnas if c in en_la_entrada}
        sin_encabezado = {i for i, c in enumerate(en_la_entrada) if not c}
        filas[nombre] = []
        for n, fila in enumerate(valores, 2):
            fila = [texto_de_celda(v) for v in fila]
            if not any(fila):
                continue
            errores += [f"{nombre}, fila {n}: la columna {get_column_letter(i + 1)} no tiene encabezado y tiene datos («{fila[i]}»)."
                        for i in sorted(sin_encabezado) if i < len(fila) and fila[i]]
            filas[nombre].append((n, {c: fila[i] if i < len(fila) else "" for c, i in posicion.items()}))
        quitadas[nombre] = [c for c in se_quitan if c in en_la_entrada]
    formulas.close()
    if errores:
        rechazar(errores)
    ignoradas = [n for n in libro.sheetnames if n not in filas]
    return filas, {n: c for n, c in quitadas.items() if c}, ignoradas


def corregir(esquema, filas, correcciones):
    """Aplica a las filas de la entrada las correcciones, {pestaña: {ID o «fila N»: {columna: texto}}}. Devuelve los
    errores de las que no se pueden aplicar y cuántas filas cambió."""
    if not (isinstance(correcciones, dict) and all(isinstance(c, dict) for c in correcciones.values())):
        return ["Las correcciones deben ser un objeto JSON {pestaña: {ID o «fila N»: {columna: texto}}}."], 0
    errores, corregidas = [], set()
    for nombre, por_fila in correcciones.items():
        if nombre not in filas:
            errores.append(f"correcciones: {nombre!r} no es una pestaña de 11/REG-ESQUEMA; opciones: "
                           f"{', '.join(filas)}.")
            continue
        p = pestana_del_esquema(esquema, nombre)
        id_ = clave(p)
        for llave, cambio in por_fila.items():
            etiqueta = f"correcciones, {nombre}, {llave}"
            numero = re.fullmatch(r"fila (\d+)", llave)
            if numero:
                elegidas = [(n, f) for n, f in filas[nombre] if n == int(numero.group(1))]
            elif id_["tipo"] != "id":
                errores.append(f"{etiqueta}: {nombre} no tiene ID propio; cada fila se corrige como «fila N».")
                continue
            else:
                elegidas = [(n, f) for n, f in filas[nombre] if f[id_["nombre"]] == llave]
            if not elegidas:
                errores.append(f"{etiqueta}: no hay una fila con ese {'número' if numero else id_['nombre']} en "
                               f"{nombre}.")
                continue
            if len(elegidas) > 1:
                errores.append(f"{etiqueta}: el ID se repite en las filas "
                               f"{', '.join(str(n) for n, _ in elegidas)}; corrige cada una como «fila N».")
                continue
            if not isinstance(cambio, dict):
                errores.append(f"{etiqueta}: debe ser un objeto {{columna: texto}}.")
                continue
            n, fila = elegidas[0]
            for columna, valor in cambio.items():
                if columna not in fila:
                    errores.append(f"{etiqueta}: {columna!r} no es una columna guardada de {nombre}.")
                elif not isinstance(valor, (str, type(None))):
                    errores.append(f"{etiqueta}: {columna} debe ser texto, entre comillas; dice {valor!r}.")
                else:
                    fila[columna] = valor or ""
                    corregidas.add((nombre, n))
    return errores, len(corregidas)


def migrar(esquema, entrada, correcciones):
    """Escribe el archivo de registros desde una hoja descargada como .xlsx o pone al día el archivo 1.x de la copia:
    sin las columnas calculadas de informes ni las retiradas, y con cada fila comprobada como en agregar después de
    aplicarle las correcciones. Con un solo error no escribe nada. Solo lee la entrada, salvo que sea el archivo
    mismo."""
    ruta = ruta_archivo(esquema)
    if not os.path.isfile(entrada):
        error(f"{entrada}: no existe el archivo que se migra.")
    if os.path.exists(ruta) and not os.path.samefile(ruta, entrada) and tiene_filas(ruta):
        error(f"{ruta} ya tiene filas; migrar no lo sobrescribe. Solo pone al día ese mismo archivo, si es de una "
              "versión anterior: `migrar` con su ruta.")
    filas, quitadas, ignoradas = leer_para_migrar(esquema, entrada)
    errores, corregidas = corregir(esquema, filas, correcciones)
    todas = {nombre: [f for _, f in de_la_pestana] for nombre, de_la_pestana in filas.items()}
    revisar = []
    for p in esquema["pestanas"]:
        id_ = clave(p)
        for n, fila in filas[p["nombre"]]:
            con_id = id_["tipo"] == "id" and fila[id_["nombre"]]
            etiqueta = f"fila {n}" + (f" ({fila[id_['nombre']]})" if con_id else "")
            revisar.append((p["nombre"], fila, etiqueta))
    errores += comprobar(esquema, todas, revisar)
    if errores:
        rechazar(errores)
    vistas = instantanea(esquema, todas)
    libro = libro_vacio(esquema)
    anexar(libro, todas)
    guardar(libro, ruta)
    guardar_instantanea(vistas)
    escritas = abrir(esquema)[1]
    return {"archivo": ruta,
            "conteos": {nombre: {"entrada": len(filas[nombre]), "archivo": len(escritas[nombre])} for nombre in filas},
            "columnas_quitadas": quitadas, "corregidas": corregidas, "pestanas_ignoradas": ignoradas}


def leer_zona_horaria():
    """La zona_horaria del bloque de 04/PREF-UBICACION, con la que se cuenta la antigüedad."""
    ruta = MODULO_04
    bloque = bloque_json(seccion(leer_modulo(ruta, "la zona horaria"), "PREF-UBICACION"))
    try:
        zona = json.loads(bloque or "null")["zona_horaria"]
    except (ValueError, TypeError, KeyError):
        zona = None
    if not isinstance(zona, str):
        error(f"{ruta}, PREF-UBICACION: falta la zona_horaria del bloque; llénala con `/configurar`, paso 4 "
              "(Preferencias).")
    try:
        return ZoneInfo(zona)
    except (ValueError, KeyError):
        error(f"{ruta}, PREF-UBICACION: zona_horaria no es una zona IANA conocida: {zona!r}. Si el nombre es "
              "correcto, al entorno virtual le falta la base de zonas horarias (en Windows: "
              "`.venv\\Scripts\\python.exe -m pip install tzdata`).")


def umbral_de_antigua():
    """Los días de la marca «antigua», de su línea en 08/EVAL-ACUERDOS; sin esa línea, los 20 por defecto."""
    ruta = MODULO_08
    acuerdos = seccion(leer_modulo(ruta, "el umbral de «antigua»"), "EVAL-ACUERDOS") or ""
    linea = re.search(r"^- \*\*Marca «antigua»:\*\*.*$", acuerdos, re.M)
    if linea is None:
        return UMBRAL_DE_ANTIGUA
    dias = re.search(r"a partir de \**(\d+)\**\s+días?\b", linea.group(0))
    if dias is None:
        error(f"{ruta}, EVAL-ACUERDOS: la línea «Marca «antigua»» debe decir «a partir de N días», con N en "
              "cifras; dice: " + linea.group(0))
    return int(dias.group(1))


def antiguedad(vacante, hoy, umbral):
    """08/EVAL-ANTIGUEDAD: «N d» desde la más antigua entre fecha_publicacion y
    primera_deteccion; «≥ N d» si solo hay primera_deteccion; «—» si además se detectó hoy."""
    publicada, detectada = vacante["fecha_publicacion"], vacante["primera_deteccion"]
    try:
        n = max((hoy - dt.datetime.strptime(f, "%Y-%m-%d").date()).days
                for f in ((detectada,) if publicada == "" else (publicada, detectada)))
    except ValueError:
        return "#VALUE!"
    if publicada == "" and n <= 0:
        return "—"
    return ("≥ " if publicada == "" else "") + f"{n} d" + (" · antigua" if n >= umbral else "")


def encaje(evaluacion):
    if evaluacion["condiciones"] == "Excluida":
        return "—"
    return evaluacion["encaje"] + (" · " + evaluacion["cobertura"] if evaluacion["cobertura"] else "")


def alertas(evaluacion):
    """Si está excluida, su detalle; si no, los núcleos insuficientes y académicos y los pendientes poco
    habituales (08/EVAL-INFORME)."""
    detalle = evaluacion["condiciones_detalle"]
    if evaluacion["condiciones"] == "Excluida":
        return "Excluida: " + detalle.replace(" | ", "; ")
    nucleos = [n for n in evaluacion["nucleos"].split(" | ") if n]

    def de_categoria(categoria):
        return "; ".join(n for n in nucleos if n.endswith(categoria)).replace(categoria, "")

    insuficientes = de_categoria(": Insuficiente")
    academicas = de_categoria(": Académica/formativa")
    habituales = ("Sueldo", "Duración", "Relación", "Horario", "Ninguno")
    pendientes = ", ".join(p for p in detalle.split(" | ") if p and p not in habituales).lower()
    partes = [titulo + texto for titulo, texto in (("Insuficiente: ", insuficientes), ("Académica: ", academicas),
                                                   ("Pendiente: ", pendientes)) if texto]
    return " · ".join(partes) or "—"


def calculadas(vacante, evaluacion, hoy, umbral):
    """Las columnas calculadas de una fila de informes, como las describe 11/REG-CAMPOS, con #N/A si su vacante
    o su evaluación no existen."""
    if vacante is None:
        de_la_vacante = dict.fromkeys(("empresa", "vacante", "antiguedad", "url_oficial"), "#N/A")
    else:
        job_id = vacante["job_id"]
        de_la_vacante = {
            "empresa": vacante["empresa_id"].replace("EMP-", ""),
            # «No publicado» sin distinguir mayúsculas.
            "vacante": ("sin Job ID" if job_id.casefold() in ("", "no publicado") else job_id) + " — "
                       + vacante["titulo"],
            "antiguedad": antiguedad(vacante, hoy, umbral),
            "url_oficial": vacante["url_oficial"]}
    if evaluacion is None:
        de_la_evaluacion = dict.fromkeys(("encaje", "alertas"), "#N/A")
    else:
        de_la_evaluacion = {"encaje": encaje(evaluacion), "alertas": alertas(evaluacion)}
    return de_la_vacante | de_la_evaluacion


def primeras(filas, columna):
    """{ID: fila}, con la primera fila de cada ID."""
    por_id = {}
    for f in filas:
        por_id.setdefault(f[columna], f)
    return por_id


def tabla_resumen(esquema, filas, busqueda_id, hoy, umbral):
    """Las filas de informes de una búsqueda, en su orden, con sus columnas calculadas a la fecha hoy."""
    vacantes = primeras(filas["vacantes"], "vacante_id")
    evaluaciones = primeras(filas["evaluaciones"], "evaluacion_id")
    columnas = [c["nombre"] for c in pestana_del_esquema(esquema, "informes")["columnas"]]
    tabla = []
    for f in filas["informes"]:
        if f["busqueda_id"] == busqueda_id:
            fila = f | calculadas(vacantes.get(f["vacante_id"]), evaluaciones.get(f["evaluacion_id"]), hoy, umbral)
            tabla.append({c: fila[c] for c in columnas})
    return tabla


def informe(esquema, busqueda_id):
    """La tabla resumen de una búsqueda (08/EVAL-INFORME): sus filas de informes, en su orden, con las
    columnas calculadas a la fecha de hoy."""
    filas = abrir(esquema)[1]
    if not any(f["busqueda_id"] == busqueda_id for f in filas["busquedas"]):
        error(f"{busqueda_id} no existe en busquedas.")
    hoy = dt.datetime.now(leer_zona_horaria()).date()
    return tabla_resumen(esquema, filas, busqueda_id, hoy, umbral_de_antigua())


def codificar(url, signos):
    """La URL con esos signos como %XX, para que no corten una liga ni una celda de tabla."""
    return re.sub(f"[{re.escape(signos)}]", lambda m: f"%{ord(m.group()):02X}", url)


def md(valor):
    """Un valor de los registros como texto de la instantánea: en una línea, con los signos de Markdown escapados
    y cada URL como autoliga, para que se lea literal y pase markdownlint. La puntuación que sigue a una URL
    queda fuera de la liga."""
    texto = []
    for i, parte in enumerate(re.split(r"(https?://\S+?)(?=[.,;:!?)»]*(?:\s|$))", " ".join(valor.split()))):
        if i % 2:
            texto.append(f"<{codificar(parte, '<>|')}>")
        else:
            texto.append(re.sub(r"([\\`*_\[\]<>|#~])", r"\\\1", parte))
    return "".join(texto)


def liga(texto, url):
    """[texto](url), o solo el texto si no hay URL."""
    if not url.startswith(("http://", "https://")):
        return md(texto)
    return f"[{md(texto)}]({codificar(url, '()<>| ')})"


def campos(fila, sin=()):
    """Las columnas con valor de una fila, una por renglón de lista."""
    return [f"- **{c}:** {md(v)}" for c, v in fila.items() if v and c not in sin]


def vista_informes(esquema, filas, encabezado, hoy, umbral):
    """informes.md: las últimas búsquedas, la más reciente primero, con su fila de busquedas y su tabla resumen."""
    busquedas = sorted(filas["busquedas"], key=lambda b: (b["fecha_hora"], b["busqueda_id"]), reverse=True)
    ultimas = busquedas[:ULTIMAS_BUSQUEDAS]
    lineas = ["# Instantánea: informes", "", encabezado, "",
              f"Búsquedas: las últimas {len(ultimas)} de {len(busquedas)}, la más reciente primero, cada una con su "
              "fila de `busquedas` y su tabla resumen (08/EVAL-INFORME)."]
    for b in ultimas:
        lineas += ["", f"## {b['busqueda_id']} · {md(b['fecha_hora'])}", ""]
        lineas += campos({c: b[c] for c in ("solicitud_alcance", "cobertura_limitaciones", "no_revisado_pendientes",
                                             "comparacion_anterior")})
        tabla = tabla_resumen(esquema, filas, b["busqueda_id"], hoy, umbral)
        if not tabla:
            lineas += ["", "Sin vacantes evaluadas en `informes`."]
            continue
        lineas += ["", "| Vacante | Encaje | Antigüedad | Alertas |", "| --- | --- | --- | --- |"]
        lineas += [f"| {liga(f['vacante'], f['url_oficial'])} | {md(f['encaje'])} | {md(f['antiguedad'])} "
                   f"| {md(f['alertas'])} |" for f in tabla]
    return "\n".join(lineas) + "\n"


def vista_vacantes(filas, encabezado, hoy, umbral):
    """vacantes.md: una sección autocontenida por vacante evaluada cuya última observación no es Cerrada."""
    observaciones = primeras(filas["observaciones"], "observacion_id")
    ultima_evaluacion = {e["vacante_id"]: e for e in filas["evaluaciones"]}
    requisitos = {}
    for r in filas["habilidades_vacante"]:
        requisitos.setdefault(r["observacion_id"], []).append(r)
    secciones = []
    for v in filas["vacantes"]:
        evaluacion = ultima_evaluacion.get(v["vacante_id"])
        observacion = observaciones.get(v["ultima_observacion_id"])
        if evaluacion is None or v["duplicado_de"] or (observacion and observacion["estado_vigencia"] == "Cerrada"):
            continue
        job_id = "sin Job ID" if v["job_id"].casefold() in ("", "no publicado") else v["job_id"]
        seccion = [f"## {v['vacante_id']} · {md(job_id)} · {v['empresa_id']} · «{md(v['titulo'])}»", ""]
        seccion += campos(v, sin=("vacante_id",)) + [f"- **antigüedad:** {md(antiguedad(v, hoy, umbral))}"]
        if observacion:
            seccion += ["", f"### Última observación: {observacion['observacion_id']}", ""]
            seccion += campos(observacion, sin=("observacion_id", "vacante_id"))
        seccion += ["", f"### Última evaluación: {evaluacion['evaluacion_id']}", ""]
        seccion += campos(evaluacion, sin=("evaluacion_id", "vacante_id"))
        # Una fila por frase literal: «Python or R» junta sus dos términos y conserva la alternativa.
        grupos = {}
        for r in requisitos.get(evaluacion["observacion_id"], []):
            normalizados, tipos = grupos.setdefault(r["habilidad_literal"], ([], []))
            for lista, valor in ((normalizados, r["habilidad_normalizada"]), (tipos, r["tipo_requisito"])):
                if valor and valor not in lista:
                    lista.append(valor)
        seccion += ["", f"### Requisitos de {evaluacion['observacion_id']}", ""]
        if grupos:
            seccion += ["| Frase literal | Normalizados | Tipo |", "| --- | --- | --- |"]
            seccion += [f"| {md(literal)} | {md('; '.join(normalizados)) or '—'} | {md('; '.join(tipos))} |"
                        for literal, (normalizados, tipos) in grupos.items()]
        else:
            seccion.append("Sin requisitos registrados.")
        secciones.append("\n".join(seccion))
    lineas = ["# Instantánea: vacantes evaluadas", "", encabezado, "",
              f"Vacantes: {len(secciones)}, con evaluación y cuya última observación no es `Cerrada`. Cada una trae sus "
              "campos de `vacantes`, su última observación, su última evaluación y los requisitos de la observación "
              "evaluada, agrupados por frase literal. No están las vacantes sin evaluación, el historial de "
              "observaciones ni `documentos`."]
    return "\n\n".join(["\n".join(lineas)] + secciones) + "\n"


def instantanea(esquema, filas):
    """Las dos vistas de la instantánea (11/REG-INSTANTANEA), {archivo: texto}, desde todas las filas. Se arman
    antes de guardar el archivo de registros: si falta lo que piden, como la zona horaria de 04, no se escribe nada."""
    ahora = dt.datetime.now(leer_zona_horaria())
    umbral = umbral_de_antigua()
    encabezado = (f"**Exportada:** {ahora:%Y-%m-%d %H:%M} ({ahora.tzinfo.key}), desde "
                  f"`{esquema['archivo']['nombre']}.xlsx`. La antigüedad está calculada a esa fecha; las demás fechas "
                  "son las de los registros. Se regenera completa con cada escritura y no se edita a mano "
                  "(11/REG-INSTANTANEA).")
    return {"informes.md": vista_informes(esquema, filas, encabezado, ahora.date(), umbral),
            "vacantes.md": vista_vacantes(filas, encabezado, ahora.date(), umbral)}


def guardar_instantanea(vistas):
    """Escribe cada vista completa en instantanea/ y avisa en stderr, sin fallar, si juntas pasan del límite."""
    carpeta = os.path.join(RAIZ, CARPETA_INSTANTANEA)
    os.makedirs(carpeta, exist_ok=True)
    tamanos = {}
    for nombre, texto in vistas.items():
        ruta = os.path.join(carpeta, nombre)
        temporal = ruta + ".tmp"
        try:
            with open(temporal, "w", encoding="utf-8", newline="\n") as f:
                f.write(texto)
            os.replace(temporal, ruta)
        finally:
            if os.path.exists(temporal):
                os.remove(temporal)
        tamanos[f"{CARPETA_INSTANTANEA}/{nombre}"] = os.path.getsize(ruta)
    total = sum(tamanos.values())
    if total > AVISO_INSTANTANEA_KB * 1024:
        print(f"Aviso: la instantánea pesa {total / 1024:.0f} KB, más de {AVISO_INSTANTANEA_KB} KB; el canal Web puede no leerla "
              "completa (11/REG-INSTANTANEA).", file=sys.stderr)
    return {"archivos": tamanos}


def leer(esquema, nombre, ids, columna):
    p = pestana_del_esquema(esquema, nombre)
    columna = columna or clave(p)["nombre"]
    if columna not in encabezados(p):
        error(f"{columna!r} no es una columna guardada de {nombre}.")
    filas = abrir(esquema)[1][nombre]
    return [f for f in filas if f[columna] in ids] if ids else filas


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ordenes = ap.add_subparsers(dest="orden", required=True)
    ordenes.add_parser("crear", help="crea el archivo desde 11/REG-ESQUEMA")
    o = ordenes.add_parser("agregar", help="agrega filas, comprobadas contra 11/REG-ESQUEMA")
    o.add_argument("filas", nargs="?", default="-", metavar="ARCHIVO.json",
                   help="{pestaña: [fila, ...]}; por defecto, de stdin")
    o = ordenes.add_parser("actualizar", help="cambia columnas de filas registradas, por su ID")
    o.add_argument("filas", nargs="?", default="-", metavar="ARCHIVO.json",
                   help="{pestaña: [{columna_id: ID, columna: valor nuevo, ...}]}; por defecto, de stdin")
    o = ordenes.add_parser("informe", help="la tabla resumen de una búsqueda, con las columnas calculadas")
    o.add_argument("busqueda_id", metavar="BUS-ID")
    o = ordenes.add_parser("leer", help="filas de una pestaña")
    o.add_argument("pestana", metavar="PESTAÑA")
    o.add_argument("--id", nargs="+", default=[], help="solo las filas con alguno de estos IDs")
    o.add_argument("--columna", help="columna donde se buscan los IDs (por defecto, la primera)")
    o = ordenes.add_parser("importar", help="registra una búsqueda desde la salida de buscar.py y las decisiones")
    o.add_argument("salida", metavar="SALIDA.json", help="la salida JSON de herramientas/buscar.py")
    o.add_argument("decisiones", metavar="DECISIONES.json",
                   help="alcance, versiones, pendientes y la decisión de cada línea «revisar»")
    o = ordenes.add_parser("migrar", help="pasa a esta versión una hoja descargada o un archivo de registros anterior")
    o.add_argument("entrada", metavar="ENTRADA.xlsx", help="la hoja descargada como .xlsx o el archivo de registros 1.x")
    o.add_argument("--correcciones", metavar="CORRECCIONES.json",
                   help='{pestaña: {ID o "fila N": {columna: texto}}}, aplicadas antes de comprobar')
    ordenes.add_parser("instantanea", help="regenera la instantánea completa que lee el canal Web")
    a = ap.parse_args()
    esquema = leer_esquema()
    if a.orden == "crear":
        res = crear(esquema)
    elif a.orden == "agregar":
        res = agregar(esquema, leer_entrada(a.filas))
    elif a.orden == "actualizar":
        res = actualizar(esquema, leer_entrada(a.filas))
    elif a.orden == "importar":
        res = importar(esquema, leer_entrada(a.salida), leer_entrada(a.decisiones))
    elif a.orden == "migrar":
        res = migrar(esquema, a.entrada, leer_entrada(a.correcciones) if a.correcciones else {})
    elif a.orden == "informe":
        res = informe(esquema, a.busqueda_id)
    elif a.orden == "instantanea":
        res = guardar_instantanea(instantanea(esquema, abrir(esquema)[1]))
    else:
        res = leer(esquema, a.pestana, a.id, a.columna)
    print(json.dumps(res, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
